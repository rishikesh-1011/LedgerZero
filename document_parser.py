"""
document_parser.py

Universal financial document and transaction statement parser.
Supports:
  - CSV / TSV / TXT (.csv, .tsv, .txt)
  - Excel (.xlsx, .xls)
  - Word Document (.docx)
  - PDF Statement (.pdf)
  - XML Statement (.xml)
  - JSON (.json)

Extracts and normalizes records into canonical transaction dicts:
  {
    "amount": float,
    "date": datetime,
    "reference_id": str,
    "vendor": str,
    "_id": str,
    "_row_index": int
  }
"""

import csv
import io
import json
import os
import re
from datetime import datetime

# Regex heuristics for column header detection
AMOUNT_COL_PATTERNS = [
    r"^amount$", r"^txn_?amount$", r"^transaction_?amount$", r"^total$",
    r"^debit$", r"^credit$", r"^dr_?cr$", r"^paid$", r"^received$", r"^val$",
    r"^value$", r"^amt$", r"^sum$", r"^net_?amount$", r"^balance$"
]

DATE_COL_PATTERNS = [
    r"^date$", r"^txn_?date$", r"^transaction_?date$", r"^value_?date$",
    r"^booking_?date$", r"^post_?date$", r"^posting_?date$", r"^entry_?date$",
    r"^settlement_?date$", r"^time$", r"^timestamp$"
]

REF_COL_PATTERNS = [
    r"^reference_?id$", r"^reference$", r"^ref$", r"^ref_?no$", r"^reference_?no$",
    r"^txn_?id$", r"^transaction_?id$", r"^utr$", r"^utr_?no$", r"^utr_?number$",
    r"^cheque_?no$", r"^chq_?no$", r"^invoice_?no$", r"^inv_?no$", r"^id$",
    r"^order_?id$", r"^payment_?id$"
]

VENDOR_COL_PATTERNS = [
    r"^vendor$", r"^vendor_?name$", r"^merchant$", r"^payee$", r"^payer$",
    r"^description$", r"^particulars$", r"^narration$", r"^party$", r"^party_?name$",
    r"^account_?name$", r"^entity$", r"^beneficiary$", r"^name$", r"^details$"
]

# Date format patterns to try when parsing string dates
DATE_FORMATS = [
    "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y", "%d/%m/%Y",
    "%Y/%m/%d", "%d.%m.%Y", "%Y.%m.%d", "%d-%b-%Y",
    "%d %b %Y", "%d %B %Y", "%b %d, %Y", "%B %d, %Y",
    "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S"
]


def _match_header(header_name, pattern_list):
    """Check if header_name matches any regex in pattern_list."""
    norm = header_name.strip().lower().replace(" ", "_")
    for pat in pattern_list:
        if re.search(pat, norm):
            return True
    return False


def _parse_amount(val):
    """Clean and parse a numeric amount from string or number."""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(abs(val))

    s = str(val).strip()
    # Remove currency symbols and formatting commas
    s = re.sub(r"[^\d.\-+]", "", s.replace(",", ""))
    if not s or s in ("-", "+", "."):
        return 0.0
    try:
        return float(abs(float(s)))
    except ValueError:
        return 0.0


def _parse_date(val):
    """Parse date from string or datetime object."""
    if val is None:
        return datetime.now()
    if isinstance(val, datetime):
        return val
    if hasattr(val, "to_pydatetime"):
        return val.to_pydatetime()

    s = str(val).strip()
    # Handle ISO / common timestamp prefixes
    s = s.split("T")[0].split(" ")[0]

    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass

    # Regex search for YYYY-MM-DD or DD-MM-YYYY
    match = re.search(r"(\d{4})[/-](\d{1,2})[/-](\d{1,2})", s)
    if match:
        try:
            return datetime(int(match.group(1)), int(match.group(2)), int(match.group(3)))
        except ValueError:
            pass

    match = re.search(r"(\d{1,2})[/-](\d{1,2})[/-](\d{4})", s)
    if match:
        try:
            return datetime(int(match.group(3)), int(match.group(2)), int(match.group(1)))
        except ValueError:
            pass

    return datetime.now()


def _detect_columns(fieldnames):
    """Map raw column headers to canonical (amount, date, reference_id, vendor) names."""
    mapping = {"amount": None, "date": None, "reference_id": None, "vendor": None}

    for col in fieldnames:
        if not mapping["amount"] and _match_header(col, AMOUNT_COL_PATTERNS):
            mapping["amount"] = col
        elif not mapping["date"] and _match_header(col, DATE_COL_PATTERNS):
            mapping["date"] = col
        elif not mapping["reference_id"] and _match_header(col, REF_COL_PATTERNS):
            mapping["reference_id"] = col
        elif not mapping["vendor"] and _match_header(col, VENDOR_COL_PATTERNS):
            mapping["vendor"] = col

    # Fallback to positional matching if some are missing
    unmapped = [c for c in fieldnames if c not in mapping.values()]
    if not mapping["amount"] and unmapped:
        mapping["amount"] = unmapped.pop(0)
    if not mapping["date"] and unmapped:
        mapping["date"] = unmapped.pop(0)
    if not mapping["vendor"] and unmapped:
        mapping["vendor"] = unmapped.pop(0)
    if not mapping["reference_id"] and unmapped:
        mapping["reference_id"] = unmapped.pop(0)

    return mapping


def _rows_from_dicts(raw_rows, source_name="upload"):
    """Normalize a list of dicts using detected column mapping."""
    if not raw_rows:
        return []

    fieldnames = list(raw_rows[0].keys())
    col_map = _detect_columns(fieldnames)

    amt_col = col_map["amount"]
    date_col = col_map["date"]
    ref_col = col_map["reference_id"]
    ven_col = col_map["vendor"]

    normalized = []
    for i, row in enumerate(raw_rows):
        amt_val = _parse_amount(row.get(amt_col)) if amt_col else 0.0
        date_val = _parse_date(row.get(date_col)) if date_col else datetime.now()
        ref_val = str(row.get(ref_col, "")).strip() if ref_col and row.get(ref_col) is not None else ""
        ven_val = str(row.get(ven_col, "Unknown")).strip() if ven_col and row.get(ven_col) is not None else "Unknown"

        # Skip completely empty or zero rows
        if amt_val <= 0.001 and not ref_val:
            continue

        normalized.append({
            "amount": amt_val,
            "date": date_val,
            "reference_id": ref_val,
            "vendor": ven_val or "Unknown",
            "_id": f"{source_name}#{i}",
            "_row_index": i
        })

    return normalized


# ---------------------------------------------------------------------------
# Format Parsers
# ---------------------------------------------------------------------------
def parse_csv(file_bytes_or_str, source_name="file.csv"):
    """Parse CSV / TSV text or bytes."""
    if isinstance(file_bytes_or_str, bytes):
        text = file_bytes_or_str.decode("utf-8", errors="replace")
    else:
        text = str(file_bytes_or_str)

    # Detect delimiter (comma, tab, semicolon, pipe)
    sample = text[:2048]
    delimiter = ","
    for d in ["\t", ";", "|", ","]:
        if d in sample and sample.count(d) > 2:
            delimiter = d
            break

    f = io.StringIO(text)
    reader = csv.DictReader(f, delimiter=delimiter)
    raw_rows = list(reader)
    return _rows_from_dicts(raw_rows, source_name)


def parse_excel(file_bytes, source_name="file.xlsx"):
    """Parse Excel spreadsheet (.xlsx, .xls) using openpyxl or pandas."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        sheet = wb.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []

        # Header in first non-empty row
        header_idx = 0
        while header_idx < len(rows) and not any(rows[header_idx]):
            header_idx += 1

        if header_idx >= len(rows):
            return []

        headers = [str(c).strip() if c is not None else f"col_{j}" for j, c in enumerate(rows[header_idx])]
        raw_rows = []
        for r in rows[header_idx + 1:]:
            if not any(r):
                continue
            row_dict = {headers[j]: r[j] for j in range(min(len(headers), len(r)))}
            raw_rows.append(row_dict)

        return _rows_from_dicts(raw_rows, source_name)
    except Exception:
        # Fallback to pandas
        import pandas as pd
        df = pd.read_excel(io.BytesIO(file_bytes))
        raw_rows = df.to_dict(orient="records")
        return _rows_from_dicts(raw_rows, source_name)


def parse_docx(file_bytes, source_name="file.docx"):
    """Parse Word document (.docx) tables or lines."""
    import docx
    doc = docx.Document(io.BytesIO(file_bytes))
    raw_rows = []

    # Check tables first
    if doc.tables:
        for table in doc.tables:
            if len(table.rows) > 1:
                headers = [cell.text.strip() for cell in table.rows[0].cells]
                for r in table.rows[1:]:
                    vals = [cell.text.strip() for cell in r.cells]
                    if any(vals):
                        row_dict = {headers[j]: vals[j] for j in range(min(len(headers), len(vals)))}
                        raw_rows.append(row_dict)

    if raw_rows:
        return _rows_from_dicts(raw_rows, source_name)

    # Fallback to paragraphs text parsing
    text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    return parse_text_lines(text, source_name)


def parse_pdf(file_bytes, source_name="file.pdf"):
    """Parse PDF statement (.pdf) extracting text and tabular transaction rows."""
    import pypdf
    reader = pypdf.PdfReader(io.BytesIO(file_bytes))
    full_text = []
    for page in reader.pages:
        t = page.extract_text()
        if t:
            full_text.append(t)

    combined_text = "\n".join(full_text)
    return parse_text_lines(combined_text, source_name)


def parse_xml(file_bytes_or_str, source_name="file.xml"):
    """Parse XML statement (.xml) extracting transaction nodes."""
    import xml.etree.ElementTree as ET

    if isinstance(file_bytes_or_str, str):
        file_bytes = file_bytes_or_str.encode("utf-8")
    else:
        file_bytes = file_bytes_or_str

    root = ET.fromstring(file_bytes)
    raw_rows = []

    # Search for repeating transaction/record/entry nodes
    entry_tags = ["transaction", "record", "entry", "row", "item", "stmtentry", "ntry", "tx", "payment"]
    found_nodes = []

    for tag in entry_tags:
        nodes = root.findall(f".//{tag}") or root.findall(f".//{tag.upper()}") or root.findall(f".//{tag.capitalize()}")
        if nodes:
            found_nodes = nodes
            break

    # If no standard tag found, find elements with child elements
    if not found_nodes:
        found_nodes = [elem for elem in root.iter() if len(elem) >= 2]

    for elem in found_nodes:
        row = {}
        # Collect direct child text or attributes
        for child in elem:
            tag_name = child.tag.split("}")[-1]  # remove namespace
            row[tag_name] = child.text.strip() if child.text else ""
        for k, v in elem.attrib.items():
            row[k] = v
        if row:
            raw_rows.append(row)

    if raw_rows:
        return _rows_from_dicts(raw_rows, source_name)

    return []


def parse_json(file_bytes_or_str, source_name="file.json"):
    """Parse JSON transaction array or nested object."""
    if isinstance(file_bytes_or_str, bytes):
        text = file_bytes_or_str.decode("utf-8", errors="replace")
    else:
        text = str(file_bytes_or_str)

    data = json.loads(text)
    if isinstance(data, list):
        return _rows_from_dicts(data, source_name)
    elif isinstance(data, dict):
        for k in ["transactions", "rows", "records", "data", "entries", "items"]:
            if k in data and isinstance(data[k], list):
                return _rows_from_dicts(data[k], source_name)
        return _rows_from_dicts([data], source_name)
    return []


def parse_text_lines(text, source_name="text"):
    """Parse raw multi-line text into transaction rows using regex patterns."""
    lines = [line.strip() for line in text.split("\n") if line.strip()]
    raw_rows = []

    # Regex for a typical bank/ledger transaction line:
    # Date (YYYY-MM-DD or DD/MM/YYYY) + Reference/UTR + Vendor + Amount
    for line in lines:
        # Find date
        date_match = re.search(r"(\d{4}[-/]\d{1,2}[-/]\d{1,2}|\d{1,2}[-/]\d{1,2}[-/]\d{2,4})", line)
        # Find amount (e.g. 1,234.56 or 50000.00)
        amt_match = re.findall(r"(?:Rs\.?|INR|\$|€|£)?\s*([0-9]{1,3}(?:,[0-9]{3})*(?:\.[0-9]{2})|[0-9]+\.[0-9]{2})", line)

        if date_match and amt_match:
            date_str = date_match.group(1)
            amt_str = amt_match[-1]  # take the last numeric currency amount

            # Reference ID match (e.g. TXN100001, REF123, UTR456)
            ref_match = re.search(r"\b(TXN\d+|REF\d+|UTR\d+|FEE\d+|INV\d+|[A-Z]{2,4}\d{5,})\b", line, re.IGNORECASE)
            ref_str = ref_match.group(1) if ref_match else ""

            # Everything else is vendor / particulars
            rem = line.replace(date_str, "").replace(amt_str, "").replace(ref_str, "")
            vendor_str = re.sub(r"[\s,\-\|]+", " ", rem).strip() or "Transaction"

            raw_rows.append({
                "amount": amt_str,
                "date": date_str,
                "reference_id": ref_str,
                "vendor": vendor_str
            })

    return _rows_from_dicts(raw_rows, source_name)


# ---------------------------------------------------------------------------
# Universal Entry Point
# ---------------------------------------------------------------------------
def parse_document(file_content, filename="statement.csv"):
    """
    Universal parser entry point. Detects format by extension / content
    and extracts normalized transaction records.

    Args:
        file_content: bytes or str
        filename: str filename with extension

    Returns:
        list of normalized transaction dicts
    """
    ext = os.path.splitext(filename)[-1].lower()

    if ext in (".csv", ".tsv", ".txt"):
        return parse_csv(file_content, filename)
    elif ext in (".xlsx", ".xls"):
        if isinstance(file_content, str):
            file_content = file_content.encode("utf-8")
        return parse_excel(file_content, filename)
    elif ext in (".docx", ".doc"):
        if isinstance(file_content, str):
            file_content = file_content.encode("utf-8")
        return parse_docx(file_content, filename)
    elif ext == ".pdf":
        if isinstance(file_content, str):
            file_content = file_content.encode("utf-8")
        return parse_pdf(file_content, filename)
    elif ext == ".xml":
        return parse_xml(file_content, filename)
    elif ext == ".json":
        return parse_json(file_content, filename)
    else:
        # Fallback: try CSV first, then text lines
        try:
            return parse_csv(file_content, filename)
        except Exception:
            text = file_content.decode("utf-8", errors="replace") if isinstance(file_content, bytes) else str(file_content)
            return parse_text_lines(text, filename)
